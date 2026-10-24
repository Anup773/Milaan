
"""
Reads a bank statement and turns each transaction into a balanced pair of source rows.

A bank statement is one-sided: it only says what moved through the bank. To fit the books it becomes two lines per
transaction, which the normal import pipeline then maps and posts like any other voucher:

    money in   (deposit)     Dr  <bank account>      Cr  <suspense account>
    money out  (withdrawal)  Dr  <suspense account>  Cr  <bank account>

The suspense account is a holding account (for example "Bank Statement Suspense"). A person later moves each amount
to the real account (rent, sales, a customer...) with an adjustment, so nothing is guessed about what a payment was for.

What is checked, and never skipped silently:
  * the running balance after every transaction (when the statement has a balance column),
  * rows that look like transactions but cannot be read (they become ERROR rows and are not posted),
  * transactions already imported from an earlier statement (marked ERROR, so they are not posted twice).

Nothing here touches the database. Duplicate checking is given the keys that already exist (see mark_duplicates).
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Set, Tuple

from ..ingestion.models import ParsedSourceRow, ParseIssue, ParseResult, Severity
from ..ingestion.money_parse import parse_amount, parse_date

ZERO = Decimal("0")
HEADER_SCAN_ROWS = 40
MAX_NAME_LENGTH = 150
# The database keeps one source row per (file, row number), so the second line of each pair is numbered
# <statement row> + this offset. Problems are always reported against the statement's own row number.
COUNTER_ROW_OFFSET = 1_000_000
SUPPORTED = (".csv", ".xlsx", ".xlsm")


class BankError(ValueError):
    """Bad input or an unreadable statement. A ValueError so the API reports it as a 400 the person can fix."""


# ---- column names ---------------------------------------------------------------------------------------------------
# Matched after normalising: lower case, anything that is not a letter or digit becomes one space.
# "Chq./Ref.No." -> "chq ref no", "Withdrawal Amt." -> "withdrawal amt". Earlier names in a list win.

ALIASES: Dict[str, Sequence[str]] = {
    "date": ("txn date", "tran date", "transaction date", "posting date", "post date", "trans date", "entry date", "date", "value date", "value dt"),
    "description": ("narration", "description", "particulars", "transaction details", "transaction remarks", "transaction description", "details", "remarks"),
    "reference": ("chq ref no", "chq ref number", "ref no cheque no", "chq no", "chqno", "cheque no", "cheque number", "instrument no", "utr no", "utr", "reference no", "ref no", "reference"),
    "withdrawal": ("withdrawal amt", "withdrawal amount", "withdrawals", "withdrawal", "debit amount", "debits", "debit", "dr amount", "dr", "paid out", "money out", "amount debited"),
    "deposit": ("deposit amt", "deposit amount", "deposits", "deposit", "credit amount", "credits", "credit", "cr amount", "cr", "paid in", "money in", "amount credited"),
    "amount": ("transaction amount", "txn amount", "amount inr", "amount"),
    "type": ("dr cr", "cr dr", "debit credit", "transaction type", "txn type", "type"),
    "balance": ("closing balance", "running balance", "available balance", "balance inr", "closing bal", "balance", "bal"),
}

_NOT_A_TRANSACTION = re.compile(
    r"^\s*(grand\s+total|total|closing\s+balance|opening\s+balance|balance|statement\s+summary|brought\s+forward|carried\s+forward|b/f|c/f|\*+)",
    re.IGNORECASE,
)


def _norm(header: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(header if header is not None else "").lower()).strip()


@dataclass
class Layout:
    columns: Dict[str, int]

    @property
    def has_two_amount_columns(self) -> bool:
        return "withdrawal" in self.columns and "deposit" in self.columns

    @property
    def is_usable(self) -> bool:
        return "date" in self.columns and (self.has_two_amount_columns or "amount" in self.columns)


def detect_layout(header_cells: Sequence[Any]) -> Layout:
    normalized = [_norm(h) for h in header_cells]
    columns: Dict[str, int] = {}
    taken: Set[int] = set()
    for field_name, names in ALIASES.items():
        for alias in names:
            hit = next((i for i, cell in enumerate(normalized) if cell == alias and i not in taken), None)
            if hit is not None:
                columns[field_name] = hit
                taken.add(hit)  # one column serves one field only
                break
    return Layout(columns)


def find_header_row(rows: Sequence[Sequence[Any]]) -> Optional[int]:
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        if detect_layout(row).is_usable:
            return index
    return None


# ---- reading single cells ---------------------------------------------------------------------------------------------

_EXTRA_DATE_FORMATS = ("%d/%m/%y", "%d-%m-%y", "%d-%b-%y", "%d %b %y", "%d %B %Y", "%d-%B-%Y", "%Y/%m/%d", "%d.%m.%y", "%b %d, %Y")
_TRAILING_TIME = re.compile(r"[ T]+\d{1,2}:\d{2}(:\d{2})?(\.\d+)?\s*([AaPp][Mm])?$")
_DR_CR_SUFFIX = re.compile(r"\s*(cr|dr|c|d)\.?$", re.IGNORECASE)


def read_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return None  # a bare number in a date column is not a date that can be trusted
    text = _TRAILING_TIME.sub("", str(value).strip())
    found = parse_date(text)  # day-first formats come before month-first ones
    if found is not None:
        return found
    for fmt in _EXTRA_DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() in ("", "-", "–", "—"))


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # a cheque number that Excel stored as 123456.0
    return str(value).strip()


def read_amount(value: Any) -> Decimal:
    """Exact decimal from a cell. '1,20,000.50', '(500)', '1,000.00 Cr' (Cr positive, Dr negative) and real numbers all work."""
    if _blank(value):
        return ZERO
    if isinstance(value, bool):
        raise ValueError(f"could not read amount: {value!r}")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))  # via text, never straight from a float's binary form
    text = str(value).strip()
    sign = 1
    suffix = _DR_CR_SUFFIX.search(text)
    if suffix and re.search(r"\d", text[: suffix.start()]):
        sign = 1 if suffix.group(1).lower() in ("cr", "c") else -1
        text = text[: suffix.start()]
    return parse_amount(text) * sign


def _debit_marker(value: Any) -> Optional[bool]:
    """True = money out, False = money in, None = not recognised."""
    text = _norm(value)
    if text in ("dr", "d", "debit", "withdrawal", "wdl", "out"):
        return True
    if text in ("cr", "c", "credit", "deposit", "dep", "in"):
        return False
    return None


def clean_account_name(value: str, label: str) -> str:
    name = (value or "").strip()
    if not name:
        raise BankError(f"{label} is required")
    if len(name) > MAX_NAME_LENGTH:
        raise BankError(f"{label} is too long (limit {MAX_NAME_LENGTH} characters)")
    if re.search(r"[\x00-\x1f]", name):
        raise BankError(f"{label} has characters that are not allowed")
    return name


# ---- the parsed statement ---------------------------------------------------------------------------------------------

@dataclass
class Transaction:
    row_number: int
    txn_date: Optional[date]
    narration: str
    reference: str
    withdrawal: Decimal
    deposit: Decimal
    balance: Optional[Decimal]
    readable: bool = True
    key: str = ""


@dataclass
class BankStats:
    transactions: int = 0
    deposits: int = 0
    withdrawals: int = 0
    total_deposits: Decimal = ZERO
    total_withdrawals: Decimal = ZERO
    opening_balance: Optional[Decimal] = None
    closing_balance: Optional[Decimal] = None
    first_date: Optional[date] = None
    last_date: Optional[date] = None
    balance_check: str = "not_checked"  # passed | failed | not_checked
    balance_mismatches: int = 0
    order: str = "oldest_first"  # oldest_first | newest_first
    skipped_rows: int = 0
    unreadable_rows: int = 0
    duplicates: int = 0

    def as_dict(self) -> dict:
        def money(v: Optional[Decimal]) -> Optional[str]:
            return None if v is None else str(v)

        return {
            "transactions": self.transactions,
            "deposits": self.deposits,
            "withdrawals": self.withdrawals,
            "total_deposits": str(self.total_deposits),
            "total_withdrawals": str(self.total_withdrawals),
            "net_movement": str(self.total_deposits - self.total_withdrawals),
            "opening_balance": money(self.opening_balance),
            "closing_balance": money(self.closing_balance),
            "first_date": self.first_date.isoformat() if self.first_date else None,
            "last_date": self.last_date.isoformat() if self.last_date else None,
            "balance_check": self.balance_check,
            "balance_mismatches": self.balance_mismatches,
            "order": self.order,
            "skipped_rows": self.skipped_rows,
            "unreadable_rows": self.unreadable_rows,
            "duplicates": self.duplicates,
        }


@dataclass
class BankParse:
    result: ParseResult
    stats: BankStats
    transactions: List[Transaction] = field(default_factory=list)


def _txn_key(bank_account: str, tx: Transaction, occurrence: int) -> str:
    material = "|".join([
        bank_account.casefold(),
        tx.txn_date.isoformat() if tx.txn_date else "",
        str(tx.withdrawal),
        str(tx.deposit),
        "" if tx.balance is None else str(tx.balance),
        _norm(tx.narration),
        _norm(tx.reference),
        str(occurrence),
    ])
    return hashlib.sha256(material.encode("utf-8")).hexdigest()


def parse_table(rows: List[List[Any]], bank_account: str, suspense_account: str, ref_tag: str = "BS", source_format: str = "bank") -> BankParse:
    """rows: the cells of the sheet from top to bottom, as read (text, or real numbers and dates)."""
    bank_account = clean_account_name(bank_account, "the bank account name")
    suspense_account = clean_account_name(suspense_account, "the suspense account name")
    if bank_account.casefold() == suspense_account.casefold():
        raise BankError("the bank account and the suspense account must be different accounts")

    header_idx = find_header_row(rows)
    if header_idx is None:
        raise BankError(
            "could not find the column headings of a bank statement in the first 40 rows. Expected a Date column, and "
            "either Withdrawal and Deposit (or Debit and Credit) columns, or a single Amount column"
        )
    layout = detect_layout(rows[header_idx])
    cols = layout.columns
    result = ParseResult(source_format=source_format, sheet_name=None, detected_columns={f: str(rows[header_idx][i]) for f, i in cols.items()})
    stats = BankStats()
    transactions: List[Transaction] = []
    lead_balance: Optional[Decimal] = None  # an "opening balance" row above the first transaction

    def cell(row: Sequence[Any], name: str) -> Any:
        idx = cols.get(name)
        return row[idx] if idx is not None and idx < len(row) else None

    last_tx: Optional[Transaction] = None
    for offset, row in enumerate(rows[header_idx + 1:], start=1):
        row_number = header_idx + offset + 1  # the row number a person sees in the file
        if all(_blank(c) for c in row):
            continue
        raw_date, raw_balance = cell(row, "date"), cell(row, "balance")
        narration, reference = _text(cell(row, "description")), _text(cell(row, "reference"))
        txn_date = read_date(raw_date)

        issues: List[ParseIssue] = []

        def amount_of(name: str, value: Any) -> Decimal:
            try:
                return read_amount(value)
            except ValueError as exc:
                issues.append(ParseIssue(row_number, name, str(exc), Severity.ERROR))
                return ZERO

        withdrawal = deposit = ZERO
        if layout.has_two_amount_columns:
            w_raw, d_raw = cell(row, "withdrawal"), cell(row, "deposit")
            amount_present = not (_blank(w_raw) and _blank(d_raw))
            withdrawal, deposit = amount_of("withdrawal", w_raw), amount_of("deposit", d_raw)
        else:
            a_raw = cell(row, "amount")
            amount_present = not _blank(a_raw)
            amount = amount_of("amount", a_raw)
            if amount_present and not issues:
                marker = _debit_marker(cell(row, "type")) if "type" in cols else None
                if "type" in cols and marker is None:
                    issues.append(ParseIssue(row_number, "type", f"could not tell whether this is money in or out from {cell(row, 'type')!r} (expected Dr or Cr)", Severity.ERROR))
                elif marker is True or (marker is None and amount < 0):
                    withdrawal = abs(amount)
                else:
                    deposit = abs(amount) if marker is False else amount

        # Rows that are not transactions: wrapped narration, opening balance, totals and footers.
        if not amount_present:
            if txn_date is None and _blank(raw_balance) and narration and last_tx is not None and not _NOT_A_TRANSACTION.match(narration):
                last_tx.narration = f"{last_tx.narration} {narration}".strip()  # a long narration wrapped onto the next row
                continue
            if txn_date is not None and not transactions and lead_balance is None and not _blank(raw_balance):
                try:
                    lead_balance = read_amount(raw_balance)
                except ValueError:
                    pass
            stats.skipped_rows += 1
            continue
        if txn_date is None and _NOT_A_TRANSACTION.match(narration or " ".join(_text(c) for c in row)):
            stats.skipped_rows += 1  # a Total / Closing Balance line
            continue

        try:
            balance = None if _blank(raw_balance) else read_amount(raw_balance)
        except ValueError:
            balance = None
            issues.append(ParseIssue(row_number, "balance", f"could not read the balance: {raw_balance!r}", Severity.WARNING))

        tx = Transaction(row_number, txn_date, narration, reference, withdrawal, deposit, balance)
        if txn_date is None:
            issues.append(ParseIssue(row_number, "date", f"unreadable or missing date: {raw_date!r}", Severity.ERROR))
        if any(i.severity == Severity.ERROR for i in issues):
            tx.readable = False
        elif withdrawal != 0 and deposit != 0:
            issues.append(ParseIssue(row_number, "withdrawal/deposit", "both a withdrawal and a deposit on one row", Severity.ERROR))
            tx.readable = False
        elif withdrawal < 0 or deposit < 0:
            issues.append(ParseIssue(row_number, "withdrawal/deposit", "a negative amount in the withdrawal or deposit column", Severity.ERROR))
            tx.readable = False
        elif withdrawal == 0 and deposit == 0:
            issues.append(ParseIssue(row_number, "withdrawal/deposit", "the amount is zero, so nothing was posted for this row", Severity.WARNING))
            stats.skipped_rows += 1
            result.issues.extend(issues)
            continue
        result.issues.extend(issues)
        transactions.append(tx)
        last_tx = tx

    if not transactions:
        raise BankError("the headings were found but there are no transactions under them")
    if transactions[-1].row_number >= COUNTER_ROW_OFFSET:
        raise BankError(f"this statement has too many rows (the limit is {COUNTER_ROW_OFFSET:,}). Download a shorter date range")

    _summarise(result, stats, transactions, lead_balance)
    # Fingerprints: an identical-looking row in the same file is a different transaction, so occurrences are counted.
    seen: Dict[str, int] = {}
    for tx in transactions:
        if tx.readable:
            base = _txn_key(bank_account, tx, 0)
            seen[base] = seen.get(base, 0) + 1
            tx.key = _txn_key(bank_account, tx, seen[base] - 1)
    for tx in transactions:
        _add_pair(result, tx, bank_account, suspense_account, ref_tag)
    return BankParse(result, stats, transactions)


def _summarise(result: ParseResult, stats: BankStats, transactions: List[Transaction], lead_balance: Optional[Decimal]) -> None:
    good = [t for t in transactions if t.readable]
    stats.unreadable_rows = len(transactions) - len(good)
    stats.transactions = len(good)
    for tx in good:
        if tx.deposit:
            stats.deposits += 1
            stats.total_deposits += tx.deposit
        else:
            stats.withdrawals += 1
            stats.total_withdrawals += tx.withdrawal
    dates = [t.txn_date for t in good if t.txn_date]
    if dates:
        stats.first_date, stats.last_date = min(dates), max(dates)

    # Statements come oldest-first or newest-first; the balance check has to walk them in true order.
    ordered = list(transactions)
    if len(good) >= 2 and good[0].txn_date and good[-1].txn_date and good[0].txn_date > good[-1].txn_date:
        stats.order = "newest_first"
        ordered.reverse()
        lead_balance = None  # a balance row above a newest-first list is the closing balance, not the opening one
    _check_balances(result, stats, ordered, lead_balance)


def _check_balances(result: ParseResult, stats: BankStats, ordered: List[Transaction], lead_balance: Optional[Decimal]) -> None:
    previous = lead_balance
    checked = mismatches = 0
    for tx in ordered:
        if not tx.readable or tx.balance is None:
            previous = None  # the chain is broken: the next row cannot be checked against this one
            continue
        if previous is not None:
            checked += 1
            expected = previous + tx.deposit - tx.withdrawal
            if expected != tx.balance:
                mismatches += 1
                result.issues.append(ParseIssue(
                    tx.row_number, "balance",
                    f"the running balance does not add up: the balance before this row plus this row should give {expected}, "
                    f"but the statement shows {tx.balance}. A row may be missing, repeated or misread above this one",
                    Severity.WARNING,
                ))
        previous = tx.balance  # carry on from what the bank says, so one gap is reported once, not on every later row

    first = next(iter(ordered), None)
    if lead_balance is not None:
        stats.opening_balance = lead_balance
    elif first is not None and first.readable and first.balance is not None:
        stats.opening_balance = first.balance - first.deposit + first.withdrawal
    last = ordered[-1] if ordered else None
    if last is not None and last.readable and last.balance is not None:
        stats.closing_balance = last.balance
    stats.balance_mismatches = mismatches
    stats.balance_check = "not_checked" if checked == 0 else ("failed" if mismatches else "passed")


def _add_pair(result: ParseResult, tx: Transaction, bank_account: str, suspense_account: str, ref_tag: str) -> None:
    """
    Two lines per transaction. Any problem on the transaction is saved on the bank line (the statement's own row
    number); that is enough to keep the whole pair from being posted, because a voucher with an error row is blocked.
    """
    reference = f"{ref_tag}-{tx.row_number}"
    description = " ".join(p for p in (tx.narration, f"[Ref {tx.reference}]" if tx.reference else "") if p).strip() or "Bank transaction"
    raw = {
        "date": tx.txn_date.isoformat() if tx.txn_date else None,
        "narration": tx.narration,
        "bank_reference": tx.reference,
        "withdrawal": str(tx.withdrawal),
        "deposit": str(tx.deposit),
        "balance": None if tx.balance is None else str(tx.balance),
        "txn_key": tx.key,
    }
    bank_debit, bank_credit = (tx.deposit, ZERO) if tx.deposit > 0 else (ZERO, tx.withdrawal)
    result.rows.append(ParsedSourceRow(tx.row_number, bank_account, tx.txn_date, description, bank_debit, bank_credit, reference, {**raw, "side": "bank"}))
    result.rows.append(ParsedSourceRow(tx.row_number + COUNTER_ROW_OFFSET, suspense_account, tx.txn_date, description, bank_credit, bank_debit, reference, {**raw, "side": "suspense"}))


# ---- duplicates -------------------------------------------------------------------------------------------------------

def keys_of(parsed: BankParse) -> List[str]:
    return [t.key for t in parsed.transactions if t.readable and t.key]


def mark_duplicates(parsed: BankParse, existing_keys: Set[str]) -> int:
    """
    Marks every transaction an earlier import already holds as an ERROR (so it is not posted twice) and returns how
    many there were. Raises BankError when every readable transaction was already imported.
    """
    if not existing_keys:
        return 0
    readable = [t for t in parsed.transactions if t.readable]
    duplicates = [t for t in readable if t.key in existing_keys]
    if duplicates and len(duplicates) == len(readable):
        raise BankError(f"this statement was already imported: all {len(readable)} transactions are in an earlier import. Nothing was saved")
    for tx in duplicates:
        parsed.result.issues.append(ParseIssue(
            tx.row_number, "duplicate",
            "this transaction is already in an earlier import (same date, amounts, balance and narration), so it will not be posted again",
            Severity.ERROR,
        ))
    parsed.stats.duplicates = len(duplicates)
    return len(duplicates)


# ---- reading files ----------------------------------------------------------------------------------------------------

def read_rows(path: Path) -> Tuple[List[List[Any]], str]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw = path.read_bytes()
        text: Optional[str] = None
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise BankError("the CSV file could not be read (unknown text encoding)")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [list(r) for r in csv.reader(io.StringIO(text, newline=""), dialect)], "bank-csv"
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise BankError(f"the Excel file could not be opened ({exc}). Is it a real .xlsx file?") from exc
        return [list(r) for r in workbook.worksheets[0].iter_rows(values_only=True)], "bank-xlsx"
    if suffix == ".xls":
        raise BankError("old .xls files cannot be read. Open it in Excel, use Save As .xlsx (or .csv), and upload that")
    raise BankError(f"unsupported file type {suffix!r}. Upload a .csv or .xlsx statement")


def parse_file(path: Path, bank_account: str, suspense_account: str, ref_tag: str = "BS") -> BankParse:
    rows, fmt = read_rows(path)
    return parse_table(rows, bank_account, suspense_account, ref_tag=ref_tag, source_format=fmt)