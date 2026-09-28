from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .detect import detect_layout, find_header_row
from .models import ParseIssue, ParseResult, Severity
from .row_parser import build_row


def _header_text(value) -> str:
    return "" if value is None else str(value)


def _no_header_issue() -> ParseIssue:
    return ParseIssue(
        row_number=0, field="header",
        message=("Could not find a header row with recognizable date/account/debit/credit "
                  "columns in the first 10 rows. Add column aliases in detect.py for this export format."),
        severity=Severity.ERROR,
    )


def parse_excel(path: Path, sheet_name: str | None = None) -> ParseResult:
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]

    rows = [row for row in ws.iter_rows(values_only=True)]
    header_texts = [[_header_text(c) for c in row] for row in rows]

    result = ParseResult(source_format="xlsx", sheet_name=ws.title, detected_columns={})

    header_idx = find_header_row(header_texts)
    if header_idx is None:
        result.issues.append(_no_header_issue())
        return result

    layout = detect_layout(header_texts[header_idx])
    result.detected_columns = {f: header_texts[header_idx][i] for f, i in layout.columns.items()}

    for offset, row in enumerate(rows[header_idx + 1:], start=1):
        row_number = header_idx + offset + 1  # +1 for 1-based, spreadsheet-visible row numbers
        if all(c is None or str(c).strip() == "" for c in row):
            continue  # blank spacer row -- common at the end of Tally exports

        def get_cell(field: str, _row=row, _cols=layout.columns):
            idx = _cols.get(field)
            return _row[idx] if idx is not None and idx < len(_row) else None

        build_row(result, row_number, get_cell, layout.columns)

    return result
